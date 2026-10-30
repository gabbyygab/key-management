import csv
import io
from datetime import timedelta
from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import Borrower, Department
from facilities.models import Building, Room
from keys.models import Key, KeyCategory
from transactions.models import BorrowTransaction, GuardShift

User = get_user_model()


class ReportsAccessTests(TestCase):
    """Verifies authentication and role-based access restrictions for reports."""

    def setUp(self):
        self.client = Client()

        self.dept = Department.objects.create(code="IT", name="Information Technology")

        # Admin user
        self.admin_user = User.objects.create_user(
            username="admin_user",
            email="admin@school.edu",
            password="adminpassword123",
            role=User.Role.ADMIN,
            approval_status=User.Approval.APPROVED,
        )

        # Superuser
        self.superuser = User.objects.create_superuser(
            username="super_user",
            email="super@school.edu",
            password="superpassword123",
            role=User.Role.FACULTY_STAFF,
            approval_status=User.Approval.APPROVED,
        )

        # Guard user
        self.guard_user = User.objects.create_user(
            username="guard_user",
            email="guard@school.edu",
            password="guardpassword123",
            role=User.Role.GUARD,
            approval_status=User.Approval.APPROVED,
        )

        # Regular user (Faculty/Staff)
        self.regular_user = User.objects.create_user(
            username="staff_user",
            email="staff@school.edu",
            password="staffpassword123",
            role=User.Role.FACULTY_STAFF,
            approval_status=User.Approval.APPROVED,
        )

    def test_anonymous_redirected_to_login(self):
        """Unauthenticated requests must be redirected to the login page."""
        response = self.client.get(reverse("reports:index"))
        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response.url)

        export_resp = self.client.get(reverse("reports:export", args=["key_utilisation"]))
        self.assertEqual(export_resp.status_code, 302)
        self.assertIn("/accounts/login/", export_resp.url)

    def test_guard_forbidden(self):
        """Guards must receive HTTP 403 Forbidden when attempting to access reports."""
        self.client.force_login(self.guard_user)
        response = self.client.get(reverse("reports:index"))
        self.assertEqual(response.status_code, 403)

        export_resp = self.client.get(reverse("reports:export", args=["key_utilisation"]))
        self.assertEqual(export_resp.status_code, 403)

    def test_regular_staff_forbidden(self):
        """Regular staff members must receive HTTP 403 Forbidden."""
        self.client.force_login(self.regular_user)
        response = self.client.get(reverse("reports:index"))
        self.assertEqual(response.status_code, 403)

    def test_admin_allowed(self):
        """ADMIN role users must be granted access (HTTP 200)."""
        self.client.force_login(self.admin_user)
        response = self.client.get(reverse("reports:index"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "reports/index.html")

    def test_superuser_allowed(self):
        """Superusers must be granted access (HTTP 200)."""
        self.client.force_login(self.superuser)
        response = self.client.get(reverse("reports:index"))
        self.assertEqual(response.status_code, 200)


class ReportsDataAndCalculationsTests(TestCase):
    """Verifies reporting metrics, business logic, overdue detection, and date filtering."""

    def setUp(self):
        self.client = Client()
        self.now = timezone.now()

        # Admin user
        self.admin = User.objects.create_user(
            username="admin1",
            email="admin1@school.edu",
            role=User.Role.ADMIN,
            approval_status=User.Approval.APPROVED,
        )
        self.client.force_login(self.admin)

        # Guards
        self.guard1 = User.objects.create_user(
            username="guard1",
            email="guard1@school.edu",
            first_name="John",
            last_name="Guard",
            role=User.Role.GUARD,
            approval_status=User.Approval.APPROVED,
        )
        self.guard2 = User.objects.create_user(
            username="guard2",
            email="guard2@school.edu",
            first_name="Jane",
            last_name="Officer",
            role=User.Role.GUARD,
            approval_status=User.Approval.APPROVED,
        )

        # Shift
        self.shift1 = GuardShift.objects.create(
            guard=self.guard1,
            post="Main Gate",
            shift_start=self.now - timedelta(days=2),
            shift_end=self.now - timedelta(days=2, hours=-8),
            is_on_duty=False,
        )

        # Building & Rooms
        self.bldg = Building.objects.create(code="ENG", name="Engineering Hall")
        self.room1 = Room.objects.create(
            building=self.bldg, room_code="101", name="Physics Lab", room_type=Room.RoomType.LABORATORY
        )
        self.room2 = Room.objects.create(
            building=self.bldg, room_code="202", name="Lecture Hall", room_type=Room.RoomType.CLASSROOM
        )

        # Category & Keys
        self.category = KeyCategory.objects.create(name="Master")
        self.key1 = Key.objects.create(
            key_code="ENG-101-A", room=self.room1, category=self.category, status=Key.Status.BORROWED
        )
        self.key2 = Key.objects.create(
            key_code="ENG-101-B", room=self.room1, category=self.category, copy_number=2, status=Key.Status.AVAILABLE
        )
        self.key3 = Key.objects.create(
            key_code="ENG-202-A", room=self.room2, category=self.category, status=Key.Status.OVERDUE
        )

        # Borrowers
        self.dept = Department.objects.create(code="CS", name="Computer Science")
        self.borrower1 = Borrower.objects.create(
            id_number="B-1001",
            first_name="Alice",
            last_name="Smith",
            borrower_type=Borrower.BorrowerType.FACULTY,
            department=self.dept,
            email="alice@school.edu",
        )
        self.borrower2 = Borrower.objects.create(
            id_number="B-1002",
            first_name="Bob",
            last_name="Jones",
            borrower_type=Borrower.BorrowerType.STUDENT,
            department=self.dept,
            email="bob@school.edu",
        )

        # Transaction 1: Returned normally 5 days ago
        self.t1 = BorrowTransaction.objects.create(
            transaction_no="TXN-20261015-0001",
            key=self.key1,
            borrower=self.borrower1,
            room=self.room1,
            borrowed_at=self.now - timedelta(days=5),
            expected_return_at=self.now - timedelta(days=5, hours=-2),
            returned_at=self.now - timedelta(days=5, hours=-1),
            released_by=self.guard1,
            received_by=self.guard2,
            return_confirmed=True,
            status=BorrowTransaction.Status.RETURNED,
        )
        BorrowTransaction.objects.filter(pk=self.t1.pk).update(borrowed_at=self.now - timedelta(days=5))

        # Transaction 2: Currently Active (borrowed yesterday, expected tomorrow - NOT overdue)
        self.t2 = BorrowTransaction.objects.create(
            transaction_no="TXN-20261019-0002",
            key=self.key1,
            borrower=self.borrower1,
            room=self.room1,
            borrowed_at=self.now - timedelta(days=1),
            expected_return_at=self.now + timedelta(days=1),
            returned_at=None,
            released_by=self.guard1,
            status=BorrowTransaction.Status.ACTIVE,
        )
        BorrowTransaction.objects.filter(pk=self.t2.pk).update(borrowed_at=self.now - timedelta(days=1))

        # Transaction 3: Currently Overdue (borrowed 3 days ago, expected 2 days ago, unreturned)
        self.t3 = BorrowTransaction.objects.create(
            transaction_no="TXN-20261017-0003",
            key=self.key3,
            borrower=self.borrower2,
            room=self.room2,
            borrowed_at=self.now - timedelta(days=3),
            expected_return_at=self.now - timedelta(days=2),
            returned_at=None,
            released_by=self.guard2,
            status=BorrowTransaction.Status.OVERDUE,
        )
        BorrowTransaction.objects.filter(pk=self.t3.pk).update(borrowed_at=self.now - timedelta(days=3))

    def test_summary_metrics(self):
        """Validates summary cards calculations across the test data."""
        response = self.client.get(reverse("reports:index"))
        self.assertEqual(response.status_code, 200)
        summary = response.context["summary"]

        self.assertEqual(summary["total_borrowings"], 3)
        self.assertEqual(summary["currently_borrowed"], 1)  # t2 is ACTIVE
        self.assertEqual(summary["overdue_count"], 1)  # t3 is OVERDUE
        self.assertEqual(summary["returned_count"], 1)  # t1 is RETURNED
        self.assertEqual(summary["active_borrowers"], 2)  # Alice and Bob
        self.assertEqual(summary["active_guards"], 2)  # guard1 and guard2

    def test_overdue_logic_does_not_flag_returned_key(self):
        """A key that has been returned must NEVER be classified as currently overdue."""
        response = self.client.get(reverse("reports:index"))
        overdue_list = response.context["overdue_list"]
        txn_numbers = [item["transaction_no"] for item in overdue_list]

        self.assertIn("TXN-20261017-0003", txn_numbers)
        self.assertNotIn("TXN-20261015-0001", txn_numbers)  # returned key
        self.assertNotIn("TXN-20261019-0002", txn_numbers)  # active key not yet due

    def test_key_utilisation_sorting(self):
        """Keys must be ordered with most frequently borrowed keys first."""
        response = self.client.get(reverse("reports:index"))
        key_utilisation = response.context["key_utilisation"]

        # key1 has 2 transactions, key3 has 1, key2 has 0
        self.assertEqual(key_utilisation[0]["key_code"], "ENG-101-A")
        self.assertEqual(key_utilisation[0]["total_borrowings"], 2)
        self.assertEqual(key_utilisation[1]["key_code"], "ENG-202-A")
        self.assertEqual(key_utilisation[1]["total_borrowings"], 1)

    def test_guard_activity_report(self):
        """Guard report must calculate releases and returns per guard."""
        response = self.client.get(reverse("reports:index"))
        guards = response.context["guard_activity"]
        guard1_data = next(g for g in guards if g["username"] == "guard1")
        guard2_data = next(g for g in guards if g["username"] == "guard2")

        # guard1 released t1 and t2
        self.assertEqual(guard1_data["keys_released"], 2)
        self.assertEqual(guard1_data["keys_returned"], 0)
        self.assertEqual(guard1_data["total_transactions"], 2)

        # guard2 released t3 and received t1 return
        self.assertEqual(guard2_data["keys_released"], 1)
        self.assertEqual(guard2_data["keys_returned"], 1)
        self.assertEqual(guard2_data["total_transactions"], 2)

    def test_room_usage_report(self):
        """Room usage must group transactions by room and avoid double counting."""
        response = self.client.get(reverse("reports:index"))
        room_usage = response.context["room_usage"]

        room1_data = next(r for r in room_usage if r["room_display"] == "ENG-101")
        room2_data = next(r for r in room_usage if r["room_display"] == "ENG-202")

        # room1 has t1 and t2
        self.assertEqual(room1_data["total_borrowings"], 2)
        self.assertEqual(room1_data["currently_borrowed"], 1)

        # room2 has t3
        self.assertEqual(room2_data["total_borrowings"], 1)
        self.assertEqual(room2_data["currently_borrowed"], 0)

    def test_borrower_frequency_report(self):
        """Borrower report must identify frequent borrowers."""
        response = self.client.get(reverse("reports:index"))
        borrowers = response.context["borrower_frequency"]

        # Alice has 2 borrowings, Bob has 1
        self.assertEqual(borrowers[0]["borrower_name"], "Alice Smith")
        self.assertEqual(borrowers[0]["total_borrowings"], 2)
        self.assertEqual(borrowers[0]["currently_active"], 1)
        self.assertEqual(borrowers[0]["returned_borrowings"], 1)

        self.assertEqual(borrowers[1]["borrower_name"], "Bob Jones")
        self.assertEqual(borrowers[1]["total_borrowings"], 1)

    def test_date_range_filtering(self):
        """Filtering by date range restricts transactions to the specified interval."""
        # Filter for the last 2 days only (should include t2 only)
        start_str = (self.now - timedelta(days=2)).strftime("%Y-%m-%d")
        end_str = self.now.strftime("%Y-%m-%d")

        response = self.client.get(
            reverse("reports:index") + f"?start_date={start_str}&end_date={end_str}"
        )
        self.assertEqual(response.status_code, 200)
        summary = response.context["summary"]

        self.assertEqual(summary["total_borrowings"], 1)  # only t2 was borrowed within last 2 days

    def test_invalid_date_range_handling(self):
        """Start date later than end date must be handled gracefully without 500 error."""
        response = self.client.get(
            reverse("reports:index") + "?start_date=2026-10-25&end_date=2026-10-10"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIsNotNone(response.context["date_error"])
        self.assertIn("Start date cannot be later than end date", response.context["date_error"])


class ReportsCSVExportTests(TestCase):
    """Verifies CSV exports, headers, data consistency, and formula injection security."""

    def setUp(self):
        self.client = Client()
        self.now = timezone.now()

        self.admin = User.objects.create_user(
            username="admin_csv",
            email="admin_csv@school.edu",
            role=User.Role.ADMIN,
            approval_status=User.Approval.APPROVED,
        )
        self.client.force_login(self.admin)

        self.guard = User.objects.create_user(
            username="guard_csv",
            email="guard_csv@school.edu",
            first_name="Sam",
            last_name="Security",
            role=User.Role.GUARD,
            approval_status=User.Approval.APPROVED,
        )

        self.bldg = Building.objects.create(code="SCI", name="Science")
        self.room = Room.objects.create(building=self.bldg, room_code="301", name="Bio Lab")
        self.cat = KeyCategory.objects.create(name="Lab")

        # Potential formula injection in key_code and borrower name
        self.key = Key.objects.create(
            key_code="=SUM(1,2)", room=self.room, category=self.cat, status=Key.Status.OVERDUE
        )
        self.borrower = Borrower.objects.create(
            id_number="@ID-9999",
            first_name="+Malicious",
            last_name="-Formula",
            borrower_type=Borrower.BorrowerType.STUDENT,
        )

        self.txn = BorrowTransaction.objects.create(
            transaction_no="TXN-CSV-001",
            key=self.key,
            borrower=self.borrower,
            room=self.room,
            borrowed_at=self.now - timedelta(days=2),
            expected_return_at=self.now - timedelta(days=1),
            returned_at=None,
            released_by=self.guard,
            status=BorrowTransaction.Status.OVERDUE,
        )

    def test_key_utilisation_csv_export(self):
        """Key utilisation CSV download generates correct content type and filename."""
        response = self.client.get(reverse("reports:export", args=["key_utilisation"]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("key_utilisation_report", response["Content-Disposition"])

        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertGreaterEqual(len(rows), 2)  # header + data
        header = rows[0]
        self.assertIn("Key Code", header)
        self.assertIn("Total Borrowings", header)

    def test_overdue_csv_export(self):
        """Overdue CSV export contains overdue record details."""
        response = self.client.get(reverse("reports:export", args=["overdue"]))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertGreaterEqual(len(rows), 2)
        header = rows[0]
        self.assertIn("Transaction No", header)
        self.assertIn("Overdue Duration", header)

    def test_guard_activity_csv_export(self):
        """Guard activity CSV export contains guard names and metrics."""
        response = self.client.get(reverse("reports:export", args=["guard_activity"]))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertGreaterEqual(len(rows), 2)
        header = rows[0]
        self.assertIn("Guard Name", header)
        self.assertIn("Keys Released", header)

    def test_room_usage_csv_export(self):
        """Room usage CSV export contains rooms and borrowing counts."""
        response = self.client.get(reverse("reports:export", args=["room_usage"]))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertGreaterEqual(len(rows), 2)
        header = rows[0]
        self.assertIn("Room", header)
        self.assertIn("Total Borrowings", header)

    def test_borrower_frequency_csv_export(self):
        """Borrower frequency CSV export contains borrower metrics."""
        response = self.client.get(reverse("reports:export", args=["borrower_frequency"]))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        self.assertGreaterEqual(len(rows), 2)
        header = rows[0]
        self.assertIn("Borrower Name", header)
        self.assertIn("Total Borrowings", header)

    def test_csv_formula_injection_defense(self):
        """Cells beginning with formula characters (=, +, -, @) must be sanitized with a leading single quote."""
        response = self.client.get(reverse("reports:export", args=["key_utilisation"]))
        content = response.content.decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)

        # Check key_code row
        key_code_found = False
        for row in rows[1:]:
            if "=SUM(1,2)" in row[0]:
                key_code_found = True
                # Must be prefixed with a single quote "'"
                self.assertTrue(
                    row[0].startswith("'="),
                    f"Formula character not neutralized: {row[0]}",
                )
        self.assertTrue(key_code_found, "Key code not found in export")

        # Check borrower frequency export for sanitized borrower names
        b_resp = self.client.get(reverse("reports:export", args=["borrower_frequency"]))
        b_content = b_resp.content.decode("utf-8")
        b_reader = csv.reader(io.StringIO(b_content))
        b_rows = list(b_reader)
        for row in b_rows[1:]:
            # First name starts with +, last name starts with -
            if "+Malicious" in row[0]:
                self.assertTrue(
                    row[0].startswith("'+"),
                    f"Formula character + not neutralized: {row[0]}",
                )
            if "@ID-9999" in row[1]:
                self.assertTrue(
                    row[1].startswith("'@"),
                    f"Formula character @ not neutralized: {row[1]}",
                )

    def test_export_via_query_param(self):
        """Exporting via ?export=<report_type> on the index view works as expected."""
        response = self.client.get(reverse("reports:index") + "?export=overdue")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("overdue_report", response["Content-Disposition"])

    def test_excel_export_format(self):
        """Excel export generates valid .xlsx workbook with correct MIME type and openpyxl structure."""
        import openpyxl
        response = self.client.get(reverse("reports:export", args=["key_utilisation"]) + "?format=xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("key_utilisation_report", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])

        # Load workbook and verify sheets and cells
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("Key Utilisation", wb.sheetnames)
        ws = wb["Key Utilisation"]
        self.assertEqual(ws["A1"].value, "KEYTRACK — UNIVERSITY KEY MANAGEMENT SYSTEM")
        # Header row at row 7
        self.assertEqual(ws["A7"].value, "Key Code")

    def test_pdf_export_format(self):
        """PDF export generates valid .pdf document with correct MIME type and PDF header bytes."""
        response = self.client.get(reverse("reports:export", args=["key_utilisation"]) + "?format=pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("key_utilisation_report", response["Content-Disposition"])
        self.assertIn(".pdf", response["Content-Disposition"])
        # Standard PDF magic header
        self.assertTrue(response.content.startswith(b"%PDF-"))

    def test_complete_all_reports_excel_export(self):
        """Complete all-reports Excel export contains Executive Summary and all 5 report sheets."""
        import openpyxl
        response = self.client.get(reverse("reports:export", args=["all"]) + "?format=xlsx")
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        expected_sheets = [
            "Executive Summary",
            "Key Utilisation",
            "Overdue Keys",
            "Guard Activity",
            "Room Usage",
            "Borrower Frequency",
        ]
        for sheet_name in expected_sheets:
            self.assertIn(sheet_name, wb.sheetnames)

    def test_complete_all_reports_pdf_export(self):
        """Complete all-reports PDF export produces valid multi-section PDF document."""
        response = self.client.get(reverse("reports:export", args=["all"]) + "?format=pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertGreater(len(response.content), 2000)

    def test_export_with_format_url_path(self):
        """URL path export_format (/reports/export/<type>/<format>/) functions identically."""
        resp_xlsx = self.client.get(reverse("reports:export_format", args=["room_usage", "xlsx"]))
        self.assertEqual(resp_xlsx.status_code, 200)
        self.assertEqual(
            resp_xlsx["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        resp_pdf = self.client.get(reverse("reports:export_format", args=["room_usage", "pdf"]))
        self.assertEqual(resp_pdf.status_code, 200)
        self.assertEqual(resp_pdf["Content-Type"], "application/pdf")
